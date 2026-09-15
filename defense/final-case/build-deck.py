"""Produce the formal six-slide local draft; placeholders are explicit."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = Path(__file__).parent
r = Presentation()
r.slide_width, r.slide_height = Inches(13.333), Inches(7.5)
BG, FG, TEAL, MUTED = '101D2B', 'F4F5F7', '4AD5C1', 'AFBDCD'

def box(slide, x,y,w,h,text,size=24,color=FG,fill=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                   Inches(w), Inches(h)) if fill else slide.shapes.add_textbox(
                                       Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.fill.background()
    tf=shape.text_frame; tf.word_wrap=True
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Noto Sans CJK SC'; p.font.size=Pt(size)
        p.font.color.rgb=RGBColor.from_string(color)
    return shape

def slide(n,title,subtitle,seconds,notes):
    s=r.slides.add_slide(r.slide_layouts[6]); s.background.fill.solid()
    s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    box(s,.65,.35,12,.4,f'QUANTUM AGENT   /   单案例教学验证                         {n:02d} / 06',13,TEAL)
    box(s,.65,1,12,1,title,34)
    box(s,.65,2,12,.7,subtitle,19,MUTED)
    box(s,.65,7,12,.3,f'正式汇报草稿 · 真实录像 / 教师结论待补       本页 {seconds} 秒',12,MUTED)
    s.notes_slide.notes_text_frame.text=notes
    return s

s=slide(1,'“衰减了，为什么还能穿过去？”','学生问题：把波函数振幅、透射概率与“完全挡住”混在一起。',40,
        '0:00–0:40。从学生困惑切入：衰减区不等于传播完全截断。展示问题，不提前给完整答案。当前案例的正式来源仍待教师核对。')
box(s,.8,3,5.7,2.7,'有限矩形势垒\nE = 5 eV，V₀ = 10 eV\na = 0.10 nm，电子质量',25,fill='203448')
box(s,7,3,5.5,2.7,'怎样判断“真的学会”？\n能解释界面条件\n能区分 r、t 与 R、T\n能迁移到 a = 0.15 nm',24)
s=slide(2,'让学生先作判断，再接受核验','教学方案：保留独立作答，把理解证据留在每一步。',45,
        '0:40–1:25。依次解释先作判断、独立尝试、证据提示、计算核验、复述迁移。不是一次问答直接给出完成标签。')
for i,t in enumerate(['先作判断','独立尝试','证据提示','计算核验','复述迁移']):
    box(s,.7+i*2.55,3.3,2.2,1.2,t,23,fill='203448')
    if i<4: box(s,2.92+i*2.55,3.6,.4,.6,'→',23,TEAL)
box(s,.8,5.1,11.8,1,'Solo 阶段保护独立作答；错误答案和 FAIL / INCONCLUSIVE 不自动升级。',23,TEAL)
s=slide(3,'真实演示：从一个判断走到一次迁移','待录制：连续产品画面，保留失败诊断；当前没有补丁版新录像。',110,
        '1:25–3:15。10秒引入，约100秒视频。画面顺序：学生判断15秒；证据与提示25秒；计算核验30秒；Teach-Back与0.15nm迁移30秒。未录制时不可用离线替身冒充。')
box(s,1.1,2.9,11.1,2.8,'▶  真实视频占位 · 约 100 秒\n待教师签核、运行版本核对及一次录制授权\n现有 DRY 原片仅为录制管线试验',26,fill='203448')
box(s,1.2,6,11,.6,'判断 15s  →  证据 25s  →  核验 30s  →  复述与迁移 30s',20,TEAL)
s=slide(4,'模型负责生成，确定性机制决定能否继续','技术机制：检索、科学核验、教学门控与录制预算分别负责一项判断。',55,
        '3:15–4:10。说明导航图谱不等于公式依据。正式来源必须经过版本、哈希、发布与条件检查。预算在每次HTTP请求前占用；原始输出保留供追查。')
for x,title,body in [(0.7,'来源','发布 / 版本 / 哈希\n原题与迁移条件绑定'),(4.95,'核验','代码实际执行指标\nT、R 精确评分'),(9.2,'教学门控','尝试 / 复述 / 迁移\n决定状态推进')]:
    box(s,x,3,3.5,2,title+'\n'+body,23,fill='203448')
box(s,.8,5.6,11.8,.9,'录制全程共用持久账本 → 每次请求先占额 → 到时取消在途任务',22,TEAL)
s=slide(5,'用冻结版本的结果说话','实测证据：本页指标从本轮最终离线日志填入，不能沿用历史通过数。',50,
        '4:10–5:00。填入本轮冻结报告中的唯一完整回归和新前端构建结果，不累加专项与完整测试。离线替身验证工程边界，不认证真实教师结论或供应商效果。')
box(s,.8,3,5.7,2.8,'冻结版指标占位\n完整 Python 离线回归：见最终报告\n前端构建 / 类型 / 密钥扫描：见最终报告\n版本：见 freeze.json',21,fill='203448')
box(s,7,3,5.5,2.8,'当前证据边界\n原两个回归：已定点复测通过\n预算：MockTransport 离线验证\n真实闭环录像：待录\n教师结论：待签核',22)
s=slide(6,'把“能回答”变成“能追查、能判定”','边界总结：本次限定为一个势垒案例及一次宽度迁移。',40,
        '5:00–5:40。总结学生价值与证据边界。实际发布、教师审批和真实录制均未在本轮执行。下一步只申请满足前置条件的一次真实运行，失败不自动重跑。')
box(s,.8,3,5.7,2.5,'已准备\n来源受控加载入口\n确定性教学与数值门控\n跨回合录制预算与离线证据',23,fill='203448')
box(s,7,3,5.5,2.5,'仍需人工完成\n教师核对原页与公式疑点\n部署镜像 / 源码 / 课程版本绑定\n确认一次真实运行申请',23)
box(s,.8,6,11.8,.5,'不把数值 PASS 扩大成整段解释正确，也不把源码测试当作部署验收。',20,TEAL)
r.save(OUT/'final-six-slides.pptx')
print(OUT/'final-six-slides.pptx')
